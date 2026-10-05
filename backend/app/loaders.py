"""Loaders that turn URLs and uploaded files into plain text chunks."""

import asyncio
import csv
import io
import ipaddress
import json
import re
import socket
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from urllib.parse import urljoin, urlparse

import httpx
import trafilatura
from langchain_text_splitters import RecursiveCharacterTextSplitter

USER_AGENT = (
    "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/124.0 Safari/537.36 AgenticRAG/0.1"
)
URL_RE = re.compile(r"https?://[^\s<>\"')\]]+", re.IGNORECASE)

TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".rst",
    ".log",
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".java",
    ".go",
    ".rs",
    ".c",
    ".cpp",
    ".h",
    ".cs",
    ".rb",
    ".php",
    ".sh",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".cfg",
    ".sql",
    ".xml",
}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {".pdf", ".docx", ".csv", ".json", ".html", ".htm"}


class LoaderError(Exception):
    pass


@dataclass
class LoadedDocument:
    title: str
    source: str
    text: str


def extract_urls(text: str) -> list[str]:
    seen: list[str] = []
    for match in URL_RE.findall(text or ""):
        url = match.rstrip(".,;:!?")
        if url not in seen:
            seen.append(url)
    return seen


def strip_urls(text: str) -> str:
    return URL_RE.sub("", text or "").strip()


def html_to_text(html: str, url: str | None = None) -> tuple[str, str]:
    """Return (title, main_text) extracted from an HTML page."""
    text = trafilatura.extract(
        html, url=url, include_comments=False, include_tables=True, favor_recall=True
    )
    meta = trafilatura.extract_metadata(html)
    title = (meta.title if meta and meta.title else None) or url or "Untitled"
    if not text:
        text = re.sub(r"<[^>]+>", " ", html)
        text = re.sub(r"\s+", " ", text)
    return title, (text or "").strip()


def _is_public_host(host: str) -> bool:
    try:
        infos = socket.getaddrinfo(host, None)
    except socket.gaierror:
        return False
    for info in infos:
        ip = ipaddress.ip_address(info[4][0].split("%")[0])
        if not ip.is_global:
            return False
    return bool(infos)


async def ensure_public_url(url: str) -> None:
    """Reject non-HTTP URLs and hosts that resolve to private/loopback/link-local addresses."""
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise LoaderError(f"Unsupported URL: {url}")
    if not await asyncio.to_thread(_is_public_host, parsed.hostname):
        raise LoaderError(f"Refusing to fetch non-public address: {url}")


async def safe_get(client: httpx.AsyncClient, url: str, max_redirects: int = 5) -> httpx.Response:
    """GET that validates the target (and every redirect hop) is a public address."""
    for _ in range(max_redirects + 1):
        await ensure_public_url(url)
        response = await client.get(url, follow_redirects=False)
        if response.is_redirect and "location" in response.headers:
            url = urljoin(str(response.url), response.headers["location"])
            continue
        return response
    raise LoaderError(f"Too many redirects for {url}")


def _pdf_filename(url: str) -> str:
    name = PurePosixPath(urlparse(url).path).name
    return name if name.lower().endswith(".pdf") else "document.pdf"


async def fetch_url(
    url: str, client: httpx.AsyncClient | None = None, timeout: float = 15.0
) -> LoadedDocument:
    owns_client = client is None
    client = client or httpx.AsyncClient(timeout=timeout, headers={"User-Agent": USER_AGENT})
    try:
        response = await safe_get(client, url)
        response.raise_for_status()
    except httpx.HTTPError as exc:
        raise LoaderError(f"Could not fetch {url}: {exc}") from exc
    finally:
        if owns_client:
            await client.aclose()

    content_type = response.headers.get("content-type", "").lower()
    if "pdf" in content_type or urlparse(url).path.lower().endswith(".pdf"):
        doc = load_file(_pdf_filename(url), response.content)
        return LoadedDocument(title=doc.title, source=url, text=doc.text)
    if "json" in content_type:
        return LoadedDocument(title=url, source=url, text=response.text)
    if "text/plain" in content_type:
        return LoadedDocument(title=url, source=url, text=response.text)

    title, text = html_to_text(response.text, url=url)
    if not text:
        raise LoaderError(f"No readable content found at {url}")
    return LoadedDocument(title=title, source=url, text=text)


def load_file(filename: str, data: bytes) -> LoadedDocument:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise LoaderError(
            f"Unsupported file type '{ext or filename}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "") for page in reader.pages]
        text = "\n\n".join(f"[Page {i + 1}]\n{p}" for i, p in enumerate(pages) if p.strip())
    elif ext == ".docx":
        from docx import Document

        document = Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(parts)
    elif ext == ".csv":
        decoded = data.decode("utf-8", errors="replace")
        rows = list(csv.reader(io.StringIO(decoded)))
        text = "\n".join(", ".join(row) for row in rows)
    elif ext == ".json":
        decoded = data.decode("utf-8", errors="replace")
        try:
            text = json.dumps(json.loads(decoded), indent=2, ensure_ascii=False)
        except json.JSONDecodeError:
            text = decoded
    elif ext in {".html", ".htm"}:
        _, text = html_to_text(data.decode("utf-8", errors="replace"))
    else:
        text = data.decode("utf-8", errors="replace")

    text = text.strip()
    if not text:
        raise LoaderError(f"No extractable text found in {filename}")
    return LoadedDocument(title=filename, source=filename, text=text)


def split_text(text: str, chunk_size: int, chunk_overlap: int) -> list[str]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=chunk_overlap)
    return [c for c in splitter.split_text(text) if c.strip()]
